# -*- coding: utf-8 -*-
"""验证三个文档的修改结果"""
import sys
from docx import Document

sys.stdout.reconfigure(encoding='utf-8')


def show_range(doc, start, end, label):
    print(f"\n--- {label} (段落 {start}-{end-1}) ---")
    for i in range(start, min(end, len(doc.paragraphs))):
        p = doc.paragraphs[i]
        style = p.style.name if p.style else "None"
        bolds = [r.bold for r in p.runs]
        print(f"[{i:3d}] <{style}> bold={bolds} {p.text[:130]}")


def find_idx(doc, sub):
    for i, p in enumerate(doc.paragraphs):
        if sub in p.text:
            return i
    return -1


# ============ 文档 1 ============
print("=" * 70)
print("验证文档 1：开发者说明书 v1.5")
print("=" * 70)
doc = Document(r"d:\family-ledger_RevB\Lunch_HP\docs\我要干饭-开发者说明书-v1.4.docx")
print(f"总段落数: {len(doc.paragraphs)}")
show_range(doc, 0, 5, "封面（前5段）")
i = find_idx(doc, "lunch_menu / lunch_orders / lunch_monthly_stats：均含 groupId")
show_range(doc, i, i + 3, "第6章 lunch_user_menu_stats 隔离说明")
i = find_idx(doc, "第9.5章")
show_range(doc, i, i + 9, "第9.5章 新章节")
i = find_idx(doc, "月度统计：lunch_monthly_stats")
show_range(doc, i, i + 3, "第11章 lunch_user_menu_stats 集合说明")

# ============ 文档 2 ============
print("\n" + "=" * 70)
print("验证文档 2：普通用户操作说明书 v1.2")
print("=" * 70)
doc = Document(r"d:\family-ledger_RevB\Lunch_HP\docs\我要干饭-普通用户操作说明书.docx")
print(f"总段落数: {len(doc.paragraphs)}")
show_range(doc, 0, 5, "封面（前5段）")
i = find_idx(doc, "菜单页从上到下依次为：")
show_range(doc, i, i + 5, "菜单页 tab 栏说明")
i = find_idx(doc, "常点列表是怎么来的")
show_range(doc, i - 1, i + 4, "Q&A 新增条目")

# ============ 文档 3 ============
print("\n" + "=" * 70)
print("验证文档 3：管理员操作说明书 v1.2")
print("=" * 70)
doc = Document(r"d:\family-ledger_RevB\Lunch_HP\docs\我要干饭-管理员操作说明书.docx")
print(f"总段落数: {len(doc.paragraphs)}")
show_range(doc, 0, 5, "封面（前5段）")
i = find_idx(doc, "常点数据需要管理员维护吗")
show_range(doc, i - 2, i + 4, "Q&A 新增条目")
