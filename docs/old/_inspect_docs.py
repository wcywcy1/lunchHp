# -*- coding: utf-8 -*-
"""探查三个 docx 文档的段落结构"""
import sys
from docx import Document

io_encoding = 'utf-8'
sys.stdout.reconfigure(encoding=io_encoding)

files = [
    r"d:\family-ledger_RevB\Lunch_HP\docs\我要干饭-开发者说明书-v1.4.docx",
    r"d:\family-ledger_RevB\Lunch_HP\docs\我要干饭-普通用户操作说明书.docx",
    r"d:\family-ledger_RevB\Lunch_HP\docs\我要干饭-管理员操作说明书.docx",
]

for f in files:
    print("=" * 80)
    print("FILE:", f)
    print("=" * 80)
    doc = Document(f)
    for i, p in enumerate(doc.paragraphs):
        text = p.text
        style = p.style.name if p.style else "None"
        # 只打印非空段落或前30段
        if text.strip() or i < 30:
            print(f"[{i:3d}] <{style}> {text[:120]}")
    print(f"\nTotal paragraphs: {len(doc.paragraphs)}\n")
