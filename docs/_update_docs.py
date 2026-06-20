# -*- coding: utf-8 -*-
"""更新 3 个 Word 文档，反映「常点」tab + lunch_user_menu_stats 代码变更"""
import copy
import sys
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

sys.stdout.reconfigure(encoding='utf-8')


def find_para(doc, substring, start=0):
    """返回第一个包含 substring 的 (index, paragraph)"""
    for i, p in enumerate(doc.paragraphs[start:], start=start):
        if substring in p.text:
            return i, p
    return -1, None


def find_last_para(doc, prefix):
    """返回最后一个以 prefix 开头的 (index, paragraph)"""
    last = (-1, None)
    for i, p in enumerate(doc.paragraphs):
        if p.text.startswith(prefix):
            last = (i, p)
    return last


def replace_in_para(para, old, new):
    """在段落的所有 run 中做字符串替换"""
    changed = False
    for r in para.runs:
        if old in r.text:
            r.text = r.text.replace(old, new)
            changed = True
    return changed


def insert_clone(source_para, ref_para, segments):
    """深拷贝 source_para 的格式，插入到 ref_para 之前。
    segments: list of (text, run_idx)，run_idx 选择拷贝哪个 source run 的 rPr。
    插入顺序：按 segments 顺序，最终排列为 [seg1, seg2, ..., ref]。
    多次对同一 ref 调用 addprevious 会按调用顺序排列。"""
    new_el = copy.deepcopy(source_para._element)
    src_runs = new_el.findall(qn('w:r'))
    templates = [copy.deepcopy(r) for r in src_runs]
    # 清除克隆段落里的 run / hyperlink
    for tag in ('w:r', 'w:hyperlink'):
        for el in new_el.findall(qn(tag)):
            new_el.remove(el)
    ref_para._element.addprevious(new_el)
    new_p = Paragraph(new_el, ref_para._parent)
    for text, run_idx in segments:
        if templates:
            idx = run_idx if 0 <= run_idx < len(templates) else 0
            new_run_el = copy.deepcopy(templates[idx])
            for t in new_run_el.findall(qn('w:t')):
                new_run_el.remove(t)
            t_el = OxmlElement('w:t')
            t_el.text = text
            t_el.set(qn('xml:space'), 'preserve')
            new_run_el.append(t_el)
            new_el.append(new_run_el)
        else:
            new_p.add_run(text)
    return new_p


# ============================================================
# 文档 1：开发者说明书 v1.4 → v1.5
# ============================================================
print("=" * 70)
print("处理文档 1：开发者说明书")
print("=" * 70)
f1 = r"d:\family-ledger_RevB\Lunch_HP\docs\我要干饭-开发者说明书-v1.4.docx"
doc = Document(f1)

# Mod1: 封面版本号 v1.4→v1.5，日期 2026-06-19→2026-06-20
idx, p_ver = find_para(doc, "v1.4")
assert p_ver is not None, "找不到封面版本行"
replace_in_para(p_ver, "v1.4", "v1.5")
replace_in_para(p_ver, "2026-06-19", "2026-06-20")
print(f"[封面版本] 段落[{idx}] -> {p_ver.text}")

# Mod2: 封面副标题追加
idx, p_sub = find_para(doc, "多租户模式切换与组织管理")
assert p_sub is not None, "找不到封面副标题"
suffix = "（v1.5：菜单页常点 tab + 个人点餐统计）"
if suffix not in p_sub.text:
    # 追加到末尾 run，保留格式
    if p_sub.runs:
        p_sub.runs[-1].text = p_sub.runs[-1].text + suffix
    else:
        p_sub.add_run(suffix)
print(f"[封面副标题] 段落[{idx}] -> {p_sub.text}")

# 定位各类源段落（用于克隆格式）
_, p_chap10 = find_para(doc, "第10章")          # 章节标题源
_, p_body51 = find_para(doc, "所有业务集合通过 groupId 字段隔离")  # Normal 正文源
_, p_bullet54 = find_para(doc, "lunch_menu / lunch_orders / lunch_monthly_stats：均含 groupId")  # List Bullet 源(单run)
_, p_bullet96 = find_para(doc, "月度统计：lunch_monthly_stats")  # List Bullet 源(双run: 粗标签+正文)

# Mod3: 第6章 —— 在 lunch_menu/lunch_orders/lunch_monthly_stats 行后补充 lunch_user_menu_stats
# 插入到 p_bullet54 的下一段之前
idx54, _ = find_para(doc, "lunch_menu / lunch_orders / lunch_monthly_stats：均含 groupId")
ref54 = doc.paragraphs[idx54 + 1]  # 下一段（云函数 lunch_menu...）
insert_clone(p_bullet54, ref54, [("lunch_user_menu_stats 同样含 groupId 字段，按组隔离。", 0)])
print(f"[第6章] 在段落[{idx54}]后插入 lunch_user_menu_stats 隔离说明")

# Mod4: 第9.5章 —— 在第10章之前插入新章节
# 重新定位第10章（上面的插入在它之前，不影响第10章本身定位）
_, p_chap10 = find_para(doc, "第10章  通用模式云函数接口")
chap95_title = "第9.5章  菜单页「常点」tab"
tab_intro = "菜单页顶部 tab 栏：全部 | 常点 | 供应商..."
b_all = "全部：所有可见菜品，按 LRU → 频率 → sortNo 排序"
b_freq = "常点：个人常点优先（userCount>0，按 lastAt 降序），不足 9 个用大众常点补齐（orderCount>0，按 lastOrderedAt 降序），平铺不分组，最多 9 个"
b_supplier = "供应商 tab：按供应商分组，tab 顺序按该供应商最近被点餐时间降序"
default_line = "默认进入「常点」tab（有记录时）。帮他人点餐时，「常点」按被帮人的 lunch_user_menu_stats 排序。"
keyfiles = "关键文件：src/hooks/useMenuFilter.ts、src/components/menu/MenuFilter.vue、src/components/menu/MenuTable.vue、wxcloud/functions/lunch_order/index.js（_upsertUserStat/getUserMenuStats/_mergePublicStats）、wxcloud/functions/lunch_menu/index.js（getMenuList 合并大众+个人统计）"

# 依次插入到第10章之前（按顺序排列）
insert_clone(p_chap10, p_chap10, [(chap95_title, 0)])      # H1 标题，克隆章节标题格式
insert_clone(p_body51, p_chap10, [(tab_intro, 0)])         # Normal 正文
insert_clone(p_bullet54, p_chap10, [(b_all, 0)])           # List Bullet
insert_clone(p_bullet54, p_chap10, [(b_freq, 0)])          # List Bullet
insert_clone(p_bullet54, p_chap10, [(b_supplier, 0)])      # List Bullet
insert_clone(p_body51, p_chap10, [(default_line, 0)])      # Normal 正文
insert_clone(p_body51, p_chap10, [(keyfiles, 0)])          # Normal 正文
print(f"[第9.5章] 在第10章前插入新章节（标题+正文+3列表项+2正文）")

# Mod5: 第11章 —— 在月度统计行后新增 lunch_user_menu_stats 行
idx96, _ = find_para(doc, "月度统计：lunch_monthly_stats")
ref96 = doc.paragraphs[idx96 + 1]  # 下一段（第12章）
new_stat_label = "个人点餐统计："
new_stat_rest = "lunch_user_menu_stats：groupId、memberId、menuId、count、lastAt（确定性 _id=groupId_memberId_menuId，支撑\"常点\"tab 个人化 LRU 排序）"
# 克隆 p_bullet96（双 run：粗标签 + 正文）
insert_clone(p_bullet96, ref96, [(new_stat_label, 0), (new_stat_rest, 1)])
print(f"[第11章] 在段落[{idx96}]后插入 lunch_user_menu_stats 集合说明")

doc.save(f1)
print("文档 1 已保存\n")


# ============================================================
# 文档 2：普通用户操作说明书 v1.1 → v1.2
# ============================================================
print("=" * 70)
print("处理文档 2：普通用户操作说明书")
print("=" * 70)
f2 = r"d:\family-ledger_RevB\Lunch_HP\docs\我要干饭-普通用户操作说明书.docx"
doc = Document(f2)

# Mod1: 封面版本号 v1.1→v1.2，日期
idx, p_ver = find_para(doc, "版本 v1.1")
assert p_ver is not None, "找不到封面版本行"
replace_in_para(p_ver, "v1.1", "v1.2")
replace_in_para(p_ver, "2026-06-19", "2026-06-20")
print(f"[封面版本] 段落[{idx}] -> {p_ver.text}")

# Mod2: 菜单页 tab 栏说明 —— 在「菜单页从上到下依次为：」后插入
idx, p_intro = find_para(doc, "菜单页从上到下依次为：")
assert p_intro is not None
# 用一个正文段落作为格式源
_, p_body_src = find_para(doc, "「我要干饭」是 HP 公司午饭点餐小程序")
if p_body_src is None:
    p_body_src = p_intro
ref = doc.paragraphs[idx + 1]  # 下一段（顶部筛选器）
tab_desc = "顶部 tab 栏：全部 | 常点 | 各供应商。「常点」会根据你的点餐记录推荐最多 9 道菜（你常点的排前面，不足时用大家常点的补齐），新菜点一次就会进常点。帮别人点餐时，常点会按那个人的记录排。进入菜单页时默认显示「常点」（有记录时）。"
insert_clone(p_body_src, ref, [(tab_desc, 0)])
print(f"[菜单页] 在段落[{idx}]后插入 tab 栏说明")

# Mod3: Q&A 末尾新增一条
last_q_idx, last_a = find_last_para(doc, "A：")
assert last_a is not None, "找不到最后一条 A："
# 找最后一条 Q 作为 Q 格式源
last_q_idx2, last_q = find_last_para(doc, "Q：")
ref_qa = doc.paragraphs[last_q_idx + 1]  # 最后 A 的下一段（空段或结尾）
new_q = "Q：常点列表是怎么来的？"
new_a = "A：根据你的点餐记录自动生成，按最近点餐时间排序（不看次数，新菜点一次就进）。不足 9 个时用大家点过的菜补齐，没被任何人点过的菜不会出现。"
insert_clone(last_q, ref_qa, [(new_q, 0)])
insert_clone(last_a, ref_qa, [(new_a, 0)])
print(f"[Q&A] 在段落[{last_q_idx}]后插入新 Q&A")

doc.save(f2)
print("文档 2 已保存\n")


# ============================================================
# 文档 3：管理员操作说明书 v1.1 → v1.2
# ============================================================
print("=" * 70)
print("处理文档 3：管理员操作说明书")
print("=" * 70)
f3 = r"d:\family-ledger_RevB\Lunch_HP\docs\我要干饭-管理员操作说明书.docx"
doc = Document(f3)

# Mod1: 封面版本号
idx, p_ver = find_para(doc, "版本 v1.1")
assert p_ver is not None, "找不到封面版本行"
replace_in_para(p_ver, "v1.1", "v1.2")
replace_in_para(p_ver, "2026-06-19", "2026-06-20")
print(f"[封面版本] 段落[{idx}] -> {p_ver.text}")

# Mod2: Q&A 末尾新增一条
last_q_idx, last_a = find_last_para(doc, "A：")
assert last_a is not None
_, last_q = find_last_para(doc, "Q：")
ref_qa = doc.paragraphs[last_q_idx + 1]
new_q = "Q：常点数据需要管理员维护吗？"
new_a = "A：不需要。「常点」tab 由系统自动统计生成（个人点餐记录 + 大众点餐记录），无需手动配置。数据存储在 lunch_user_menu_stats 集合，按组隔离。"
insert_clone(last_q, ref_qa, [(new_q, 0)])
insert_clone(last_a, ref_qa, [(new_a, 0)])
print(f"[Q&A] 在段落[{last_q_idx}]后插入新 Q&A")

doc.save(f3)
print("文档 3 已保存\n")

print("全部 3 个文档处理完成。")
