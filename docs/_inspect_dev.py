# -*- coding: utf-8 -*-
"""探查开发者说明书前 50 段"""
import sys
from docx import Document

sys.stdout.reconfigure(encoding='utf-8')

f = r"d:\family-ledger_RevB\Lunch_HP\docs\我要干饭-开发者说明书-v1.4.docx"
doc = Document(f)
print(f"Total paragraphs: {len(doc.paragraphs)}\n")
for i, p in enumerate(doc.paragraphs[:50]):
    text = p.text
    style = p.style.name if p.style else "None"
    print(f"[{i:3d}] <{style}> {text[:150]}")
