# -*- coding: utf-8 -*-
"""更新 3 个 Word 文档，反映首页去金额改点单数 + 菜单页去价格代码变更"""
import copy
import sys
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

sys.stdout.reconfigure(encoding='utf-8')


def find_para(doc, substring, start=0):
    """返回第一个包含 substring 的 (index, paragraph)"""
    for i, p in enumerate(doc.paragraphs[start:], start=start):
        if substring in p.text:
            return i, p
    return -1, None


def replace_in_para(para, old, new):
    """在段落的所有 run 中做字符串替换"""
    changed = False
    for r in para.runs:
        if old in r.text:
            r.text = r.text.replace(old, new)
            changed = True
    return changed


# ============================================================
# 文档 1：开发者说明书 v1.5 → v1.6
# ============================================================
print("=" * 70)
print("处理文档 1：开发者说明书")
print("=" * 70)
f1 = r"d:\family-ledger_RevB\Lunch_HP\docs\我要干饭-开发者说明书-v1.4.docx"
doc = Document(f1)

# Mod1: 封面版本号 v1.5→v1.6（版本行含 "·  2026-06-20"，与副标题行区分）
idx, p_ver = find_para(doc, "v1.5  \u00b7  2026-06-20")
if p_ver is None:
    idx, p_ver = find_para(doc, "v1.5")
assert p_ver is not None, "找不到封面版本行"
replace_in_para(p_ver, "v1.5", "v1.6")
print(f"[封面版本] 段落[{idx}] -> {p_ver.text}")

# Mod2: 封面副标题更新
idx, p_sub = find_para(doc, "多租户模式切换与组织管理")
assert p_sub is not None, "找不到封面副标题"
old_suffix = "（v1.5：菜单页常点 tab + 个人点餐统计）"
new_suffix = "（v1.6：菜单页常点 tab + 个人点餐统计 + 首页/菜单页去金额改点单数）"
if old_suffix in p_sub.text:
    replace_in_para(p_sub, old_suffix, new_suffix)
elif "v1.5" in p_sub.text:
    replace_in_para(p_sub, "v1.5", "v1.6")
print(f"[封面副标题] 段落[{idx}] -> {p_sub.text}")

doc.save(f1)
print("文档 1 已保存\n")


# ============================================================
# 文档 2：普通用户操作说明书 v1.2 → v1.3
# ============================================================
print("=" * 70)
print("处理文档 2：普通用户操作说明书")
print("=" * 70)
f2 = r"d:\family-ledger_RevB\Lunch_HP\docs\我要干饭-普通用户操作说明书.docx"
doc = Document(f2)

# Mod1: 封面版本号 v1.2→v1.3
idx, p_ver = find_para(doc, "版本 v1.2")
assert p_ver is not None, "找不到封面版本行"
replace_in_para(p_ver, "v1.2", "v1.3")
print(f"[封面版本] 段落[{idx}] -> {p_ver.text}")

# Mod2: 月度卡片描述 —— 金额改为点单数
idx, p_card = find_para(doc, "月度卡片：月合计金额")
assert p_card is not None, "找不到月度卡片描述"
replace_in_para(p_card, "月合计金额（橙色）、今日金额（蓝色）", "月点单数（橙色）、今日点单数（蓝色）")
print(f"[月度卡片] 段落[{idx}] -> {p_card.text}")

# Mod3: 菜单列表列字段 —— 去掉序号和价格
idx, p_menu = find_para(doc, "菜单列表：仅显示")
assert p_menu is not None, "找不到菜单列表描述"
replace_in_para(p_menu, "（序号、供应商、餐品、价格）", "（餐品名、供应商）")
print(f"[菜单列表] 段落[{idx}] -> {p_menu.text}")

doc.save(f2)
print("文档 2 已保存\n")


# ============================================================
# 文档 3：管理员操作说明书 v1.2 → v1.3
# ============================================================
print("=" * 70)
print("处理文档 3：管理员操作说明书")
print("=" * 70)
f3 = r"d:\family-ledger_RevB\Lunch_HP\docs\我要干饭-管理员操作说明书.docx"
doc = Document(f3)

# Mod1: 封面版本号 v1.2→v1.3
idx, p_ver = find_para(doc, "版本 v1.2")
assert p_ver is not None, "找不到封面版本行"
replace_in_para(p_ver, "v1.2", "v1.3")
print(f"[封面版本] 段落[{idx}] -> {p_ver.text}")

doc.save(f3)
print("文档 3 已保存\n")

print("全部 3 个文档处理完成。")
